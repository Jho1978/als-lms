import os
import random
import re
import secrets
import string
import io
import uuid
from functools import wraps
from datetime import datetime

from flask import Flask, render_template, redirect, url_for, flash, request, abort, send_file
from flask_login import (
    login_user, logout_user, login_required, current_user,
)
from sqlalchemy.exc import IntegrityError
from sqlalchemy import or_

from config import Config
from extensions import db, login_manager, csrf
from models import (
    User, Subject, Enrollment, Question, ExamAttempt, ExamAnswer,
    ROLE_ADMIN, ROLE_TEACHER, ROLE_STUDENT, LEVEL_CHOICES, LEVEL_LABELS,
    LEVEL_JUNIOR_HIGH,
)
from forms import (
    LoginForm, CreateUserForm, SubjectForm, UploadQuestionsForm, QuestionForm,
    BulkAddStudentsForm,
)


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)

    os.makedirs(os.path.join(app.instance_path), exist_ok=True)
    os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)
    os.makedirs(app.config["QUESTION_IMAGE_FOLDER"], exist_ok=True)

    db.init_app(app)
    login_manager.init_app(app)
    csrf.init_app(app)

    with app.app_context():
        db.create_all()
        _run_lightweight_migrations()
        _seed_default_admin()

    @app.context_processor
    def inject_now():
        return {"now": datetime.utcnow()}

    register_routes(app)
    return app


def _run_lightweight_migrations():
    """Add newly-introduced columns to an already-existing SQLite database.

    This project doesn't use Alembic, so on upgrade an older lms.db won't
    have brand-new columns yet. This adds them in place, without touching
    existing data, so people don't have to delete their database to update.
    """
    from sqlalchemy import inspect, text

    inspector = inspect(db.engine)
    existing_tables = inspector.get_table_names()

    if "subjects" in existing_tables:
        subject_columns = {c["name"] for c in inspector.get_columns("subjects")}
        if "level" not in subject_columns:
            db.session.execute(
                text(f"ALTER TABLE subjects ADD COLUMN level VARCHAR(30) DEFAULT '{LEVEL_JUNIOR_HIGH}'")
            )
            db.session.commit()

    if "users" in existing_tables:
        user_columns = {c["name"] for c in inspector.get_columns("users")}
        if "grade_level" not in user_columns:
            db.session.execute(text("ALTER TABLE users ADD COLUMN grade_level VARCHAR(30)"))
            db.session.commit()

    if "questions" in existing_tables:
        question_columns = {c["name"] for c in inspector.get_columns("questions")}
        if "image_filename" not in question_columns:
            db.session.execute(text("ALTER TABLE questions ADD COLUMN image_filename VARCHAR(255)"))
            db.session.commit()


def _seed_default_admin():
    """Create a default admin account on first run so the app is usable immediately."""
    if not User.query.filter_by(role=ROLE_ADMIN).first():
        admin = User(
            full_name="System Administrator",
            username="admin",
            email="admin@example.com",
            role=ROLE_ADMIN,
        )
        admin.set_password("admin123")
        db.session.add(admin)
        db.session.commit()
        print(">>> Default admin created — username: admin / password: admin123 (change this!)")


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


def role_required(*roles):
    def decorator(view_func):
        @wraps(view_func)
        def wrapped(*args, **kwargs):
            if not current_user.is_authenticated:
                return login_manager.unauthorized()
            if current_user.role not in roles:
                abort(403)
            return view_func(*args, **kwargs)
        return wrapped
    return decorator


def register_routes(app):

    # ---------------------------------------------------------------- MAIN
    @app.route("/")
    def index():
        if current_user.is_authenticated:
            return redirect(url_for(f"{current_user.role}_dashboard"))
        subject_count = Subject.query.count()
        teacher_count = User.query.filter_by(role=ROLE_TEACHER).count()
        student_count = User.query.filter_by(role=ROLE_STUDENT).count()
        return render_template(
            "index.html",
            subject_count=subject_count,
            teacher_count=teacher_count,
            student_count=student_count,
        )

    @app.errorhandler(403)
    def forbidden(e):
        return render_template("error.html", code=403, message="You don't have permission to view this page."), 403

    @app.errorhandler(404)
    def not_found(e):
        return render_template("error.html", code=404, message="Page not found."), 404

    # ---------------------------------------------------------------- AUTH
    @app.route("/login", methods=["GET", "POST"])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for(f"{current_user.role}_dashboard"))
        form = LoginForm()
        if form.validate_on_submit():
            user = User.query.filter_by(username=form.username.data.strip()).first()
            if user and user.check_password(form.password.data):
                login_user(user)
                flash(f"Welcome back, {user.full_name}!", "success")
                return redirect(url_for(f"{user.role}_dashboard"))
            flash("Invalid username or password.", "danger")
        return render_template("login.html", form=form)

    @app.route("/register", methods=["GET", "POST"])
    def register():
        """Public self-registration — students only."""
        form = CreateUserForm()
        if form.validate_on_submit():
            username = form.username.data.strip()
            email = (form.email.data or "").strip() or None
            if User.query.filter_by(username=username).first():
                flash("That username is already taken.", "danger")
                return render_template("register.html", form=form)
            if email and User.query.filter_by(email=email).first():
                flash("An account with that email already exists. Try logging in instead.", "danger")
                return render_template("register.html", form=form)
            student = User(
                full_name=form.full_name.data.strip(),
                username=username,
                email=email,
                role=ROLE_STUDENT,
                grade_level=form.grade_level.data or None,
            )
            student.set_password(form.password.data)
            db.session.add(student)
            try:
                db.session.commit()
            except IntegrityError:
                db.session.rollback()
                flash("That username or email is already registered.", "danger")
                return render_template("register.html", form=form)
            flash("Account created! You can now log in.", "success")
            return redirect(url_for("login"))
        return render_template("register.html", form=form)

    @app.route("/logout")
    @login_required
    def logout():
        logout_user()
        flash("You have been logged out.", "info")
        return redirect(url_for("index"))

    # --------------------------------------------------------------- ADMIN
    @app.route("/admin/dashboard")
    @role_required(ROLE_ADMIN)
    def admin_dashboard():
        return render_template(
            "admin/dashboard.html",
            teacher_count=User.query.filter_by(role=ROLE_TEACHER).count(),
            student_count=User.query.filter_by(role=ROLE_STUDENT).count(),
            subject_count=Subject.query.count(),
            attempt_count=ExamAttempt.query.count(),
        )

    @app.route("/admin/teachers", methods=["GET", "POST"])
    @role_required(ROLE_ADMIN)
    def admin_teachers():
        form = CreateUserForm()
        if form.validate_on_submit():
            username = form.username.data.strip()
            email = (form.email.data or "").strip() or None
            if User.query.filter_by(username=username).first():
                flash("That username is already taken.", "danger")
            elif email and User.query.filter_by(email=email).first():
                flash("An account with that email already exists.", "danger")
            else:
                teacher = User(
                    full_name=form.full_name.data.strip(),
                    username=username,
                    email=email,
                    role=ROLE_TEACHER,
                )
                teacher.set_password(form.password.data)
                db.session.add(teacher)
                try:
                    db.session.commit()
                    flash(f"Teacher account created for {teacher.full_name}.", "success")
                    return redirect(url_for("admin_teachers"))
                except IntegrityError:
                    db.session.rollback()
                    flash("That username or email is already registered.", "danger")
        teachers = User.query.filter_by(role=ROLE_TEACHER).order_by(User.full_name).all()
        return render_template("admin/teachers.html", form=form, teachers=teachers)

    @app.route("/admin/students", methods=["GET", "POST"])
    @role_required(ROLE_ADMIN)
    def admin_students():
        form = CreateUserForm()
        if form.validate_on_submit():
            username = form.username.data.strip()
            email = (form.email.data or "").strip() or None
            if User.query.filter_by(username=username).first():
                flash("That username is already taken.", "danger")
            elif email and User.query.filter_by(email=email).first():
                flash("An account with that email already exists.", "danger")
            else:
                student = User(
                    full_name=form.full_name.data.strip(),
                    username=username,
                    email=email,
                    role=ROLE_STUDENT,
                    grade_level=form.grade_level.data or None,
                )
                student.set_password(form.password.data)
                db.session.add(student)
                try:
                    db.session.commit()
                    flash(f"Student account created for {student.full_name}.", "success")
                    return redirect(url_for("admin_students"))
                except IntegrityError:
                    db.session.rollback()
                    flash("That username or email is already registered.", "danger")
        students = User.query.filter_by(role=ROLE_STUDENT).order_by(User.full_name).all()
        return render_template("admin/students.html", form=form, students=students)

    @app.route("/students/bulk-add", methods=["GET", "POST"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def bulk_add_students():
        # Teachers only see/enroll into their own subjects; admins see all.
        if current_user.is_admin():
            subject_choices = Subject.query.order_by(Subject.name).all()
        else:
            subject_choices = (
                Subject.query.filter_by(teacher_id=current_user.id).order_by(Subject.name).all()
            )

        form = BulkAddStudentsForm()
        form.subject_id.choices = [(0, "-- Don't enroll right now --")] + [
            (s.id, f"{s.name} ({s.level_label()})") for s in subject_choices
        ]

        preselect_subject_id = request.args.get("subject_id", type=int)
        if request.method == "GET" and preselect_subject_id:
            valid_ids = {s.id for s in subject_choices}
            if preselect_subject_id in valid_ids:
                form.subject_id.data = preselect_subject_id

        results = None
        if form.validate_on_submit():
            target_subject = None
            if form.subject_id.data:
                target_subject = next(
                    (s for s in subject_choices if s.id == form.subject_id.data), None
                )
                if target_subject is None:
                    abort(403)  # tried to enroll into a subject they don't own

            grade_level = form.grade_level.data or None
            raw_lines = [ln.strip() for ln in form.names.data.splitlines()]
            names = [ln for ln in raw_lines if ln]

            results = []
            taken_usernames = set()
            created_count = 0

            for name in names:
                # Basic sanity check — reject lines that look like stray punctuation
                # or accidental pasted junk rather than a name.
                if len(name) < 2 or not re.search(r"[A-Za-z]", name):
                    results.append({"name": name, "status": "skipped", "reason": "Doesn't look like a valid name"})
                    continue

                username = _generate_username(name, taken_usernames)
                taken_usernames.add(username)
                temp_password = _generate_temp_password()

                student = User(
                    full_name=name,
                    username=username,
                    role=ROLE_STUDENT,
                    grade_level=grade_level,
                )
                student.set_password(temp_password)
                db.session.add(student)
                try:
                    db.session.commit()
                except IntegrityError:
                    db.session.rollback()
                    results.append({"name": name, "status": "skipped", "reason": "Could not create (duplicate?)"})
                    continue

                if target_subject:
                    db.session.add(Enrollment(student_id=student.id, subject_id=target_subject.id))
                    db.session.commit()

                results.append({
                    "name": name,
                    "username": username,
                    "password": temp_password,
                    "status": "created",
                })
                created_count += 1

            if target_subject:
                flash(
                    f"Created {created_count} student account(s) and enrolled them in {target_subject.name}.",
                    "success",
                )
            else:
                flash(f"Created {created_count} student account(s).", "success")

        return render_template(
            "teacher/bulk_add_students.html",
            form=form,
            results=results,
        )

    @app.route("/admin/subjects", methods=["GET", "POST"])
    @role_required(ROLE_ADMIN)
    def admin_subjects():
        form = SubjectForm()
        form.teacher_id.choices = [(0, "-- Unassigned --")] + [
            (t.id, t.full_name) for t in User.query.filter_by(role=ROLE_TEACHER).all()
        ]
        if form.validate_on_submit():
            subject = Subject(
                name=form.name.data.strip(),
                description=(form.description.data or "").strip(),
                level=form.level.data,
                teacher_id=form.teacher_id.data or None,
                passing_score=form.passing_score.data,
                time_limit_minutes=form.time_limit_minutes.data,
            )
            db.session.add(subject)
            db.session.commit()
            flash(f"Subject '{subject.name}' created.", "success")
            return redirect(url_for("admin_subjects"))
        subjects = Subject.query.order_by(Subject.name).all()
        teachers = User.query.filter_by(role=ROLE_TEACHER).order_by(User.full_name).all()
        return render_template("admin/subjects.html", form=form, subjects=subjects, teachers=teachers)

    @app.route("/admin/subjects/<int:subject_id>/assign", methods=["POST"])
    @role_required(ROLE_ADMIN)
    def admin_assign_teacher(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        teacher_id = request.form.get("teacher_id", type=int)
        subject.teacher_id = teacher_id or None
        db.session.commit()
        flash(f"Teacher assignment updated for {subject.name}.", "success")
        return redirect(url_for("admin_subjects"))

    @app.route("/admin/results")
    @role_required(ROLE_ADMIN)
    def admin_results():
        attempts = ExamAttempt.query.order_by(ExamAttempt.taken_at.desc()).all()
        return render_template("admin/results.html", attempts=attempts)

    @app.route("/admin/results/export.xlsx")
    @role_required(ROLE_ADMIN)
    def admin_results_export():
        attempts = ExamAttempt.query.order_by(ExamAttempt.taken_at.desc()).all()
        buffer = _build_results_workbook(
            attempts, title="Campus LMS — All Exam Results", include_subject_column=True
        )
        filename = f"all_exam_results_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # ------------------------------------------------------------- TEACHER
    @app.route("/teacher/dashboard")
    @role_required(ROLE_TEACHER)
    def teacher_dashboard():
        subjects = Subject.query.filter_by(teacher_id=current_user.id).order_by(Subject.name).all()
        return render_template("teacher/dashboard.html", subjects=subjects)

    @app.route("/teacher/subjects/create", methods=["GET", "POST"])
    @role_required(ROLE_TEACHER)
    def teacher_create_subject():
        form = SubjectForm()
        del form.teacher_id  # teacher creating for themself, no need to pick
        if form.validate_on_submit():
            subject = Subject(
                name=form.name.data.strip(),
                description=(form.description.data or "").strip(),
                level=form.level.data,
                teacher_id=current_user.id,
                passing_score=form.passing_score.data,
                time_limit_minutes=form.time_limit_minutes.data,
            )
            db.session.add(subject)
            db.session.commit()
            flash(f"Subject '{subject.name}' created. Now build your test questionnaire.", "success")
            return redirect(url_for("teacher_manage_questions", subject_id=subject.id))
        return render_template("teacher/create_subject.html", form=form)

    @app.route("/teacher/subjects/<int:subject_id>/questions", methods=["GET"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def teacher_manage_questions(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        question_form = QuestionForm()
        upload_form = UploadQuestionsForm()
        questions = (
            Question.query.filter_by(subject_id=subject.id).order_by(Question.id).all()
        )
        return render_template(
            "teacher/manage_questions.html",
            subject=subject,
            question_form=question_form,
            upload_form=upload_form,
            questions=questions,
            questions_per_exam=app.config["QUESTIONS_PER_EXAM"],
        )

    @app.route("/teacher/subjects/<int:subject_id>/questions/add", methods=["POST"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def teacher_add_question(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        question_form = QuestionForm()
        if question_form.validate_on_submit():
            image_filename = _save_question_image(
                question_form.image.data,
                app.config["QUESTION_IMAGE_FOLDER"],
                app.config["ALLOWED_IMAGE_EXTENSIONS"],
            )
            db.session.add(
                Question(
                    subject_id=subject.id,
                    question_text=question_form.question_text.data.strip(),
                    choice_a=question_form.choice_a.data.strip(),
                    choice_b=question_form.choice_b.data.strip(),
                    choice_c=question_form.choice_c.data.strip(),
                    choice_d=question_form.choice_d.data.strip(),
                    correct_answer=question_form.correct_answer.data,
                    image_filename=image_filename,
                )
            )
            db.session.commit()
            remaining = max(app.config["QUESTIONS_PER_EXAM"] - subject.question_count(), 0)
            flash(
                f"Question added! ({subject.question_count()} total"
                + (f" — {remaining} more recommended to reach a full "
                   f"{app.config['QUESTIONS_PER_EXAM']}-question test)" if remaining else ")"),
                "success",
            )
        else:
            for field_errors in question_form.errors.values():
                for err in field_errors:
                    flash(err, "danger")
        return redirect(url_for("teacher_manage_questions", subject_id=subject.id))

    @app.route("/teacher/subjects/<int:subject_id>/questions/upload", methods=["POST"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def teacher_upload_questions(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        upload_form = UploadQuestionsForm()
        if upload_form.validate_on_submit():
            f = upload_form.file.data
            filename = f.filename.lower()
            try:
                added, skipped = _import_questions(subject, f, filename)
                db.session.commit()
                flash(f"Imported {added} question(s). {skipped} row(s) skipped (incomplete).", "success")
            except Exception as exc:  # noqa: BLE001
                flash(f"Could not import file: {exc}", "danger")
        else:
            for field_errors in upload_form.errors.values():
                for err in field_errors:
                    flash(err, "danger")
        return redirect(url_for("teacher_manage_questions", subject_id=subject.id))

    @app.route("/teacher/subjects/<int:subject_id>/questions/<int:question_id>/delete", methods=["POST"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def teacher_delete_question(subject_id, question_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        question = Question.query.get_or_404(question_id)
        if question.subject_id != subject.id:
            abort(404)
        if question.image_filename:
            _delete_question_image(question.image_filename, app.config["QUESTION_IMAGE_FOLDER"])
        db.session.delete(question)
        db.session.commit()
        flash("Question removed.", "info")
        return redirect(url_for("teacher_manage_questions", subject_id=subject.id))

    @app.route("/teacher/subjects/<int:subject_id>/results")
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def teacher_subject_results(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        attempts = (
            ExamAttempt.query.filter_by(subject_id=subject.id)
            .order_by(ExamAttempt.taken_at.desc())
            .all()
        )
        return render_template("teacher/results.html", subject=subject, attempts=attempts)

    @app.route("/teacher/subjects/<int:subject_id>/results/export.xlsx")
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def teacher_subject_results_export(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        attempts = (
            ExamAttempt.query.filter_by(subject_id=subject.id)
            .order_by(ExamAttempt.taken_at.desc())
            .all()
        )
        buffer = _build_results_workbook(
            attempts, title=f"Exam Results — {subject.name}", include_subject_column=False
        )
        safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", subject.name).strip("_") or "subject"
        filename = f"{safe_name}_results_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # ------------------------------------------------ SUBJECT SETTINGS / TIMER
    @app.route("/subjects/<int:subject_id>/edit", methods=["GET", "POST"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def subject_edit(subject_id):
        """Edit a subject's settings, including resetting/changing its exam
        timer (time_limit_minutes) and passing score."""
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)

        form = SubjectForm(obj=subject)
        if current_user.is_admin():
            form.teacher_id.choices = [(0, "-- Unassigned --")] + [
                (t.id, t.full_name) for t in User.query.filter_by(role=ROLE_TEACHER).all()
            ]
            if request.method == "GET":
                form.teacher_id.data = subject.teacher_id or 0
        else:
            del form.teacher_id  # teachers can't reassign ownership

        if form.validate_on_submit():
            old_limit = subject.time_limit_minutes
            subject.name = form.name.data.strip()
            subject.description = (form.description.data or "").strip()
            subject.level = form.level.data
            subject.passing_score = form.passing_score.data
            subject.time_limit_minutes = form.time_limit_minutes.data
            if current_user.is_admin():
                subject.teacher_id = form.teacher_id.data or None
            db.session.commit()

            if old_limit != subject.time_limit_minutes:
                flash(
                    f"Saved. Exam timer updated to {subject.time_limit_minutes} minutes — "
                    "this applies the next time a student starts (or retakes) the exam.",
                    "success",
                )
            else:
                flash("Subject settings saved.", "success")
            return redirect(
                url_for("admin_subjects") if current_user.is_admin() else url_for("teacher_dashboard")
            )
        return render_template("teacher/edit_subject.html", form=form, subject=subject)

    # ------------------------------------------------ MANUAL ENROLLMENT
    @app.route("/subjects/<int:subject_id>/students", methods=["GET"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def subject_manage_enrollment(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)

        level_filter = request.args.get("level", "").strip()
        search = request.args.get("q", "").strip()

        query = User.query.filter_by(role=ROLE_STUDENT)
        if level_filter:
            query = query.filter_by(grade_level=level_filter)
        if search:
            like = f"%{search}%"
            query = query.filter(
                or_(User.full_name.ilike(like), User.username.ilike(like))
            )
        students = query.order_by(User.full_name).all()

        enrolled_ids = {
            e.student_id for e in Enrollment.query.filter_by(subject_id=subject.id).all()
        }

        return render_template(
            "teacher/manage_enrollment.html",
            subject=subject,
            students=students,
            enrolled_ids=enrolled_ids,
            level_choices=LEVEL_CHOICES,
            level_filter=level_filter,
            search=search,
            enrolled_count=len(enrolled_ids),
        )

    @app.route("/subjects/<int:subject_id>/students/enroll", methods=["POST"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def subject_enroll_student(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        student_id = request.form.get("student_id", type=int)
        student = User.query.filter_by(id=student_id, role=ROLE_STUDENT).first_or_404()
        already = Enrollment.query.filter_by(student_id=student.id, subject_id=subject.id).first()
        if already:
            flash(f"{student.full_name} is already enrolled.", "warning")
        else:
            db.session.add(Enrollment(student_id=student.id, subject_id=subject.id))
            db.session.commit()
            flash(f"Enrolled {student.full_name} in {subject.name}.", "success")
        return redirect(
            url_for(
                "subject_manage_enrollment",
                subject_id=subject.id,
                level=request.args.get("level", ""),
                q=request.args.get("q", ""),
            )
        )

    @app.route("/subjects/<int:subject_id>/students/unenroll", methods=["POST"])
    @role_required(ROLE_TEACHER, ROLE_ADMIN)
    def subject_unenroll_student(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if subject.teacher_id != current_user.id and not current_user.is_admin():
            abort(403)
        student_id = request.form.get("student_id", type=int)
        enrollment = Enrollment.query.filter_by(student_id=student_id, subject_id=subject.id).first()
        if enrollment:
            db.session.delete(enrollment)
            db.session.commit()
            flash("Student removed from subject.", "info")
        return redirect(
            url_for(
                "subject_manage_enrollment",
                subject_id=subject.id,
                level=request.args.get("level", ""),
                q=request.args.get("q", ""),
            )
        )

    # ------------------------------------------------------------- STUDENT
    @app.route("/student/dashboard")
    @role_required(ROLE_STUDENT)
    def student_dashboard():
        enrollments = Enrollment.query.filter_by(student_id=current_user.id).all()
        enrolled_subject_ids = {e.subject_id for e in enrollments}
        subjects = Subject.query.filter(Subject.id.in_(enrolled_subject_ids)).all() if enrolled_subject_ids else []
        recent_attempts = (
            ExamAttempt.query.filter_by(student_id=current_user.id)
            .order_by(ExamAttempt.taken_at.desc())
            .limit(5)
            .all()
        )
        return render_template("student/dashboard.html", subjects=subjects, recent_attempts=recent_attempts)

    @app.route("/student/subjects/<int:subject_id>/exam", methods=["GET", "POST"])
    @role_required(ROLE_STUDENT)
    def student_take_exam(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        enrolled = Enrollment.query.filter_by(
            student_id=current_user.id, subject_id=subject.id
        ).first()
        if not enrolled:
            flash("You're not enrolled in this subject. Ask your teacher or the administrator to enroll you.", "warning")
            return redirect(url_for("student_dashboard"))

        if not subject.questions:
            flash("No questions have been uploaded for this subject yet.", "info")
            return redirect(url_for("student_dashboard"))

        if request.method == "POST":
            question_ids = [int(v) for v in request.form.getlist("question_ids")]
            questions = Question.query.filter(Question.id.in_(question_ids)).all()
            questions_by_id = {q.id: q for q in questions}

            score = 0
            answers_to_save = []
            for qid in question_ids:
                q = questions_by_id.get(qid)
                if not q:
                    continue
                selected = request.form.get(f"answer_{qid}")
                is_correct = bool(selected) and selected.upper() == q.correct_answer.upper()
                if is_correct:
                    score += 1
                answers_to_save.append(
                    ExamAnswer(question_id=qid, selected_answer=selected, is_correct=is_correct)
                )

            total = len(question_ids)
            percentage = round((score / total) * 100, 2) if total else 0.0
            passed = percentage >= subject.passing_score

            attempt = ExamAttempt(
                student_id=current_user.id,
                subject_id=subject.id,
                score=score,
                total=total,
                percentage=percentage,
                passed=passed,
            )
            attempt.answers = answers_to_save
            db.session.add(attempt)
            db.session.commit()
            return redirect(url_for("student_result_detail", attempt_id=attempt.id))

        # GET: build the exam — pull up to QUESTIONS_PER_EXAM questions, shuffled,
        # and shuffle each question's displayed choice order too.
        pool = list(subject.questions)
        random.shuffle(pool)
        exam_questions = pool[: app.config["QUESTIONS_PER_EXAM"]]

        display_questions = []
        for q in exam_questions:
            options = list(q.choices().items())  # [('A', text), ('B', text), ...]
            random.shuffle(options)
            display_questions.append({"question": q, "options": options})

        return render_template(
            "student/take_exam.html", subject=subject, display_questions=display_questions
        )

    @app.route("/student/results")
    @role_required(ROLE_STUDENT)
    def student_results():
        attempts = (
            ExamAttempt.query.filter_by(student_id=current_user.id)
            .order_by(ExamAttempt.taken_at.desc())
            .all()
        )
        return render_template("student/results.html", attempts=attempts)

    @app.route("/student/results/export.xlsx")
    @role_required(ROLE_STUDENT)
    def student_results_export():
        attempts = (
            ExamAttempt.query.filter_by(student_id=current_user.id)
            .order_by(ExamAttempt.taken_at.desc())
            .all()
        )
        buffer = _build_results_workbook(
            attempts, title=f"Exam Results — {current_user.full_name}", include_subject_column=True
        )
        filename = f"my_exam_results_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/student/results/<int:attempt_id>")
    @role_required(ROLE_STUDENT)
    def student_result_detail(attempt_id):
        attempt = ExamAttempt.query.get_or_404(attempt_id)
        if attempt.student_id != current_user.id:
            abort(403)
        return render_template("student/result_detail.html", attempt=attempt)


def _build_results_workbook(attempts, title, include_subject_column=False):
    """Build an .xlsx workbook of exam results (score, percentage, letter
    grade, PASS/FAIL) from a list of ExamAttempt rows, returned as an
    in-memory BytesIO ready to send as a file download."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Exam Results"

    ws.merge_cells("A1:H1" if include_subject_column else "A1:G1")
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="left")

    header = ["#", "Student Name", "Username"]
    if include_subject_column:
        header.append("Subject")
    header += ["Score", "Total", "Percentage", "Grade", "Result", "Date Taken"]

    header_row = 3
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, label in enumerate(header, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row_num = header_row + 1
    for i, attempt in enumerate(attempts, start=1):
        col = 1
        ws.cell(row=row_num, column=col, value=i); col += 1
        ws.cell(row=row_num, column=col, value=attempt.student.full_name); col += 1
        ws.cell(row=row_num, column=col, value=attempt.student.username); col += 1
        if include_subject_column:
            ws.cell(row=row_num, column=col, value=attempt.subject.name); col += 1
        ws.cell(row=row_num, column=col, value=attempt.score); col += 1
        ws.cell(row=row_num, column=col, value=attempt.total); col += 1
        pct_cell = ws.cell(row=row_num, column=col, value=attempt.percentage / 100); col += 1
        pct_cell.number_format = "0.00%"
        ws.cell(row=row_num, column=col, value=attempt.letter_grade()); col += 1
        result_cell = ws.cell(row=row_num, column=col, value="PASSED" if attempt.passed else "FAILED"); col += 1
        result_cell.font = Font(
            bold=True, color="198754" if attempt.passed else "DC3545"
        )
        ws.cell(row=row_num, column=col, value=attempt.taken_at.strftime("%Y-%m-%d %H:%M"))
        row_num += 1

    # Reasonable column widths
    widths = [4, 24, 16] + ([20] if include_subject_column else []) + [8, 8, 12, 8, 10, 18]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _save_question_image(file_storage, upload_folder, allowed_extensions):
    """Save an uploaded question image under a random filename and return
    just the filename (relative to static/question_images/), or None if no
    file was provided."""
    if not file_storage or not getattr(file_storage, "filename", ""):
        return None
    original = file_storage.filename
    ext = original.rsplit(".", 1)[-1].lower() if "." in original else ""
    if ext not in allowed_extensions:
        return None
    unique_name = f"{uuid.uuid4().hex}.{ext}"
    file_storage.save(os.path.join(upload_folder, unique_name))
    return unique_name


def _delete_question_image(filename, upload_folder):
    try:
        path = os.path.join(upload_folder, filename)
        if os.path.isfile(path):
            os.remove(path)
    except OSError:
        pass


def _generate_username(full_name, taken_usernames):
    """Build a predictable username from a full name (first initial + last
    name), automatically appending a number if it's already taken — either
    in the database or elsewhere in the current batch."""
    parts = [p for p in re.split(r"\s+", full_name.strip()) if p]
    if not parts:
        base = "student"
    elif len(parts) == 1:
        base = parts[0].lower()
    else:
        base = (parts[0][0] + parts[-1]).lower()

    base = re.sub(r"[^a-z0-9]", "", base) or "student"
    base = base[:60]  # leave room for a numeric suffix under the 80-char column limit

    candidate = base
    suffix = 1
    while candidate in taken_usernames or User.query.filter_by(username=candidate).first():
        suffix += 1
        candidate = f"{base}{suffix}"
    return candidate


def _generate_temp_password(length=8):
    """A short, random temporary password, avoiding easily-confused
    characters (0/O, 1/l/I) so it's easy to read off a printed slip."""
    alphabet = "abcdefghjkmnpqrstuvwxyzACDEFGHJKMNPQRSTUVWXYZ23456789"
    return "".join(secrets.choice(alphabet) for _ in range(length))


def _import_questions(subject, file_storage, filename):
    """Parse an uploaded Excel/CSV questionnaire and create Question rows.

    Expected columns (case-insensitive): Question, A, B, C, D, Answer
    Uses openpyxl / the stdlib csv module directly (no pandas) to avoid
    pulling in a heavy dependency that needs a C/Fortran build toolchain
    on some Windows Python versions.
    """
    import csv
    import io

    def _normalize_header(cells):
        return [str(c if c is not None else "").strip().lower() for c in cells]

    rows = []  # list of dicts keyed by normalized header
    if filename.endswith(".csv"):
        raw = file_storage.read()
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if not all_rows:
            raise ValueError("The file is empty.")
        header = _normalize_header(all_rows[0])
        for data_row in all_rows[1:]:
            if not any(str(v).strip() for v in data_row):
                continue  # skip blank lines
            padded = list(data_row) + [""] * (len(header) - len(data_row))
            rows.append(dict(zip(header, padded)))
    else:
        import openpyxl

        wb = openpyxl.load_workbook(file_storage, read_only=True, data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            raise ValueError("The file is empty.")
        header = _normalize_header(all_rows[0])
        for data_row in all_rows[1:]:
            if data_row is None or not any(
                (v is not None and str(v).strip()) for v in data_row
            ):
                continue  # skip blank rows
            padded = list(data_row) + [None] * (len(header) - len(data_row))
            rows.append(dict(zip(header, padded)))

    required = {"question", "a", "b", "c", "d", "answer"}
    missing = required - set(header)
    if missing:
        raise ValueError(
            f"Missing required column(s): {', '.join(sorted(missing))}. "
            "Expected headers: Question, A, B, C, D, Answer"
        )

    added, skipped = 0, 0
    for row in rows:
        q_text = str(row.get("question") or "").strip()
        a = str(row.get("a") or "").strip()
        b = str(row.get("b") or "").strip()
        c = str(row.get("c") or "").strip()
        d = str(row.get("d") or "").strip()
        answer = str(row.get("answer") or "").strip().upper()

        if not q_text or not a or not b or not c or not d or answer not in {"A", "B", "C", "D"}:
            skipped += 1
            continue

        db.session.add(
            Question(
                subject_id=subject.id,
                question_text=q_text,
                choice_a=a,
                choice_b=b,
                choice_c=c,
                choice_d=d,
                correct_answer=answer,
            )
        )
        added += 1

    return added, skipped


app = create_app()

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
